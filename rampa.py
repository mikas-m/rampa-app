import re
import io
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title="Raspodjela tura po rampama", page_icon="🚚")
st.title("🚚 Raspodjela tura po rampama")
st.write("Uploadaj Excel file s turama, a aplikacija će ih automatski rasporediti na rampe.")

# ==== PODESIVE POSTAVKE (u sidebaru) ====
st.sidebar.header("Postavke")
BROJ_RAMPI = st.sidebar.number_input("Broj rampi", min_value=1, max_value=20, value=5)
MAX_TURA_PO_RAMPI = st.sidebar.number_input("Max tura po rampi", min_value=1, max_value=10, value=2)
TEZINA_UDIO = st.sidebar.slider(
    "Udio težine u balansiranju (0 = samo broj narudžbi, 1 = samo težina)",
    min_value=0.0, max_value=1.0, value=0.5, step=0.1
)
trazeni_brojevi_input = st.sidebar.text_input(
    "Brojevi u nazivu lista koji označavaju smjenu (odvojeno zarezom)",
    value="6, 10, 14, 16"
)
TRAZENI_BROJEVI = [b.strip() for b in trazeni_brojevi_input.split(",") if b.strip()]

uploaded_file = st.file_uploader("Uploadaj Excel file (.xlsx)", type=["xlsx"])


# ==== POMOĆNE FUNKCIJE ====

def sheet_odgovara(sheet_name, brojevi):
    for broj in brojevi:
        pattern = r'(?<!\d)' + re.escape(broj) + r'(?!\d)'
        if re.search(pattern, sheet_name):
            return True
    return False


def pronadji_header_redak(raw_df):
    for i, row in raw_df.iterrows():
        if row.astype(str).str.contains("Tour", case=False, na=False).any():
            return i
    return None


def pronadji_kolonu(kolone, kljucna_rijec):
    for kol in kolone:
        if kljucna_rijec.lower() in kol.lower():
            return kol
    return None


def obradi_list(file_bytes, sheet_name, broj_rampi, max_tura_po_rampi, tezina_udio, log):
    raw = pd.read_excel(file_bytes, sheet_name=sheet_name, header=None)
    header_idx = pronadji_header_redak(raw)

    if header_idx is None:
        log.append(f"⏭️ Preskačem '{sheet_name}': nema kolone 'Tour', vjerojatno nije relevantan list.")
        return None

    file_bytes.seek(0)
    df = pd.read_excel(file_bytes, sheet_name=sheet_name, header=header_idx)
    df.columns = df.columns.astype(str).str.strip()

    kol_isell = pronadji_kolonu(df.columns, "iSell")
    kol_gew = pronadji_kolonu(df.columns, "Gew")
    kol_tour = pronadji_kolonu(df.columns, "Tour")

    nedostaje = [naziv for naziv, kol in
                 [("iSell", kol_isell), ("Gew", kol_gew), ("Tour", kol_tour)]
                 if kol is None]

    if nedostaje:
        log.append(f"⏭️ Preskačem '{sheet_name}': nedostaju kolone koje sadrže {nedostaje}. "
                    f"Pronađene kolone: {df.columns.tolist()}")
        return None

    df = df.rename(columns={kol_isell: "iSell", kol_gew: "Gew", kol_tour: "Tour"})
    df = df.dropna(subset=["Tour"]).copy()

    if df.empty:
        log.append(f"⏭️ Preskačem '{sheet_name}': nema podataka nakon čišćenja.")
        return None

    # brojevi tura kao tekst, bez decimala
    df["Tour"] = df["Tour"].astype(int).astype(str)

    tour_summary = (
        df.groupby("Tour")
          .agg(broj_narudzbi=("iSell", "count"), ukupna_tezina=("Gew", "sum"))
          .reset_index()
    )

    broj_tura = len(tour_summary)
    max_tura_po_rampi_lokalno = max_tura_po_rampi
    if broj_tura > broj_rampi * max_tura_po_rampi:
        max_tura_po_rampi_lokalno = -(-broj_tura // broj_rampi)  # ceil
        log.append(f"⚠️ '{sheet_name}': {broj_tura} tura, povećavam limit na "
                    f"{max_tura_po_rampi_lokalno} tura/rampi.")

    max_broj = tour_summary["broj_narudzbi"].max()
    max_tez = tour_summary["ukupna_tezina"].max()

    def skor(broj, tez):
        b = broj / max_broj if max_broj else 0
        t = tez / max_tez if max_tez else 0
        return tezina_udio * t + (1 - tezina_udio) * b

    tour_summary["skor"] = tour_summary.apply(
        lambda r: skor(r["broj_narudzbi"], r["ukupna_tezina"]), axis=1
    )
    tour_summary = tour_summary.sort_values("skor", ascending=False).reset_index(drop=True)

    rampe = {
        i: {"ture": [], "broj_narudzbi": 0, "ukupna_tezina": 0.0}
        for i in range(1, broj_rampi + 1)
    }

    for _, red in tour_summary.iterrows():
        kandidati = [r for r in rampe if len(rampe[r]["ture"]) < max_tura_po_rampi_lokalno]
        if not kandidati:
            kandidati = list(rampe.keys())

        najbolja_rampa = min(
            kandidati,
            key=lambda r: skor(rampe[r]["broj_narudzbi"], rampe[r]["ukupna_tezina"])
        )

        rampe[najbolja_rampa]["ture"].append(red["Tour"])
        rampe[najbolja_rampa]["broj_narudzbi"] += red["broj_narudzbi"]
        rampe[najbolja_rampa]["ukupna_tezina"] += red["ukupna_tezina"]

    # preslagivanje: najviše narudžbi (pa najveća težina) -> Rampa 1, najmanje -> zadnja rampa
    poredak = sorted(
        rampe.items(),
        key=lambda item: (item[1]["broj_narudzbi"], item[1]["ukupna_tezina"]),
        reverse=True
    )

    rampe_preslozene = {}
    for novi_broj, (_, podaci) in enumerate(poredak, start=1):
        rampe_preslozene[novi_broj] = podaci

    log.append(f"✅ '{sheet_name}' obrađen ({broj_tura} tura).")
    return rampe_preslozene


def obradi_excel(uploaded_file, broj_rampi, max_tura_po_rampi, tezina_udio, trazeni_brojevi):
    log = []
    file_bytes = io.BytesIO(uploaded_file.getvalue())
    xls = pd.ExcelFile(file_bytes)

    listovi = [s for s in xls.sheet_names if sheet_odgovara(s, trazeni_brojevi)]
    log.append(f"Pronađeni listovi za obradu: {listovi}")

    rezultati = {}
    for sheet in listovi:
        file_bytes.seek(0)
        rampe = obradi_list(file_bytes, sheet, broj_rampi, max_tura_po_rampi, tezina_udio, log)
        if rampe is not None:
            rezultati[sheet] = rampe

    # spremi rezultate natrag u file (u memoriji), na nove listove
    output = io.BytesIO(uploaded_file.getvalue())
    wb = load_workbook(output)

    for sheet, rampe in rezultati.items():
        novi_naziv = f"{sheet} - Rampa"[:31]
        if novi_naziv in wb.sheetnames:
            del wb[novi_naziv]
        ws = wb.create_sheet(novi_naziv)
        ws.append([sheet.upper()])
        for r in range(1, broj_rampi + 1):
            ture_str = ", ".join(str(t) for t in rampe[r]["ture"])
            ws.append([f"RAMPA {r} -> {ture_str}"])

    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)

    return final_output, rezultati, log


# ==== GLAVNI DIO SUČELJA ====

if uploaded_file is not None:
    if st.button("Obradi file"):
        with st.spinner("Obrađujem..."):
            try:
                rezultat_file, rezultati, log = obradi_excel(
                    uploaded_file, BROJ_RAMPI, MAX_TURA_PO_RAMPI, TEZINA_UDIO, TRAZENI_BROJEVI
                )
            except Exception as e:
                st.error(f"Došlo je do greške: {e}")
                st.stop()

        st.subheader("Log obrade")
        for linija in log:
            st.write(linija)

        if rezultati:
            st.subheader("Pregled rezultata")
            for sheet, rampe in rezultati.items():
                st.markdown(f"**{sheet.upper()}**")
                for r in range(1, BROJ_RAMPI + 1):
                    ture_str = ", ".join(str(t) for t in rampe[r]["ture"])
                    st.write(f"RAMPA {r} -> {ture_str}")

            st.download_button(
                label="📥 Preuzmi obrađeni Excel",
                data=rezultat_file,
                file_name="rezultat_rampe.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            st.warning("Nijedan list nije uspješno obrađen. Provjeri log iznad.")
else:
    st.info("Čekam da uploadaš Excel file.")
